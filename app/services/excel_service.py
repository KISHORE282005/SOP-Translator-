"""
Excel Processing Service for PO Translation System.

Reads and analyzes workbooks with openpyxl. The translated workbook is
written by the sheet copier, which keeps the template of the sheets intact.
"""

import re
from pathlib import Path
from typing import Any, Dict, List, Set

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook as OpenpyxlWorkbook
from openpyxl.worksheet.worksheet import Worksheet

from app.services.sheet_copier import SheetCopy, write_sheet_copies
from app.utils.logger import setup_logger

logger = setup_logger(__name__)

# Excel limits sheet titles to 31 characters
MAX_SHEET_TITLE_LENGTH = 31


class ExcelService:
    """
    Service for reading, analyzing, and generating translated Excel workbooks.
    """

    PO_PATTERN = re.compile(r'^PO[-/]\d{4}[-/]?\d*$', re.IGNORECASE)
    MATERIAL_PATTERN = re.compile(r'^(MAT|ABC|HYD)[-_]?\d+$', re.IGNORECASE)
    PART_PATTERN = re.compile(r'^(PART|JCB)[-_]?\d+$', re.IGNORECASE)
    EMPLOYEE_PATTERN = re.compile(r'^EMP\d+$', re.IGNORECASE)
    VENDOR_PATTERN = re.compile(r'^V\d+$', re.IGNORECASE)
    NUMBER_PATTERN = re.compile(r'^[+-]?\d[\d,]*(\.\d+)?%?$')
    DATE_PATTERN = re.compile(r'^\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}$|^\d{4}-\d{2}-\d{2}$')
    URL_PATTERN = re.compile(r'^(https?://|www\.)', re.IGNORECASE)
    EMAIL_PATTERN = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')
    CURRENCY_PATTERN = re.compile(r'^[$₹£€]\s*\d[\d,]*(\.\d+)?$')
    # Devanagari (Hindi) or Tamil letters
    TRANSLATED_PATTERN = re.compile(r'[ऀ-ॿ஀-௿]')

    # Punctuation that may surround a code inside a sentence, e.g. "(MAT-1001),"
    TOKEN_PUNCTUATION = '.,;:!?()[]{}"\''

    def read_workbook(self, file_path: Path) -> OpenpyxlWorkbook:
        """Read an Excel workbook from file path."""
        logger.info("Workbook opened", filename=str(file_path))
        workbook = load_workbook(file_path, data_only=False)
        sheet_names = [sheet.title for sheet in workbook.worksheets]
        logger.info("Sheets detected", sheet_count=len(sheet_names), sheet_names=sheet_names)
        return workbook

    def identify_translatable_content(self, workbook: OpenpyxlWorkbook) -> Dict[str, Any]:
        """Analyze workbook and identify cells containing translatable English text."""
        english_texts: Dict[str, None] = {}  # dict keeps first-seen order and removes duplicates
        cell_mapping = []

        for sheet in self._source_sheets(workbook):
            for row in sheet.iter_rows():
                for cell in row:
                    # Only plain text cells are translated. Numbers, dates and
                    # booleans are stored as non-string values and stay as they are.
                    if not isinstance(cell.value, str):
                        continue
                    if cell.data_type == 'f' or cell.value.startswith('='):
                        continue

                    value = cell.value
                    if not value.strip():
                        continue
                    if self._is_protected_value(value.strip()):
                        continue

                    translatable_parts = self._extract_translatable_parts(value)
                    if translatable_parts:
                        for part in translatable_parts:
                            english_texts[part] = None
                        cell_mapping.append({
                            "sheet_name": sheet.title,
                            "row": cell.row,
                            "column": cell.column,
                            "original_value": value,
                            "translatable_parts": translatable_parts,
                            "is_mixed": self._is_mixed_content(value)
                        })

        logger.info("Translatable content identified",
                    total_texts=len(english_texts), total_cells=len(cell_mapping))
        return {"english_texts": list(english_texts), "cell_mapping": cell_mapping}

    def _source_sheets(self, workbook: OpenpyxlWorkbook) -> List[Worksheet]:
        """Sheets to translate: the ones that are not already in Tamil or Hindi."""
        return [
            sheet for sheet in workbook.worksheets
            if not any(
                isinstance(cell.value, str) and self.TRANSLATED_PATTERN.search(cell.value)
                for row in sheet.iter_rows() for cell in row
            )
        ]

    def _is_protected_value(self, value: str) -> bool:
        """Check if a whole cell value should NOT be translated (protected value)."""
        if self._is_code(value): return True
        if self.NUMBER_PATTERN.match(value): return True
        if self.DATE_PATTERN.match(value): return True
        if self.CURRENCY_PATTERN.match(value): return True
        if value.startswith('='): return True
        return False

    def _is_code(self, value: str) -> bool:
        """Check if a value is an identifier that must be kept exactly as it is."""
        if self.PO_PATTERN.match(value): return True
        if self.MATERIAL_PATTERN.match(value): return True
        if self.PART_PATTERN.match(value): return True
        if self.EMPLOYEE_PATTERN.match(value): return True
        if self.VENDOR_PATTERN.match(value): return True
        if self.URL_PATTERN.match(value): return True
        if self.EMAIL_PATTERN.match(value): return True
        # Any single word mixing letters and digits is a code: JCB-12345, M8x40, X200.
        # So is a part name written with underscores: ARM_SUPPORT_PLATE
        if not re.search(r'\s', value):
            if re.search(r'[A-Za-z]', value) and re.search(r'[\d_]', value):
                return True
        return False

    def _is_code_token(self, token: str) -> bool:
        """Check if a word inside a sentence is a protected code."""
        clean_token = token.strip(self.TOKEN_PUNCTUATION)
        if not clean_token:
            return False
        return self._is_code(clean_token)

    def _extract_translatable_parts(self, value: str) -> List[str]:
        """
        Extract translatable text from a cell value, preserving protected values.

        The value is cut at line breaks and at protected codes, so
        "MAT-1001 Hydraulic Pump" gives ["Hydraulic Pump"]. Each part is an exact
        substring of the cell value. Plain numbers stay inside their part
        ("Net 30 Days") because the translator needs them for context.
        """
        translatable_parts = []

        for line in value.splitlines():
            segment_start = None
            segment_end = None

            for match in re.finditer(r'\S+', line):
                if self._is_code_token(match.group()):
                    if segment_start is not None:
                        translatable_parts.append(line[segment_start:segment_end])
                        segment_start = None
                    continue
                if segment_start is None:
                    segment_start = match.start()
                segment_end = match.end()

            if segment_start is not None:
                translatable_parts.append(line[segment_start:segment_end])

        return [part for part in translatable_parts if self._looks_like_text(part)]

    def _looks_like_text(self, value: str) -> bool:
        """Check if a value looks like natural language (English) text."""
        alpha_count = len(re.findall(r'[A-Za-z]', value))
        if alpha_count < 2:
            return False
        total_count = len(re.sub(r'\s', '', value))
        return alpha_count / total_count > 0.5

    def _is_mixed_content(self, value: str) -> bool:
        """Check if a cell value contains both codes and text."""
        has_code = False
        has_text = False
        for token in value.split():
            if self._is_code_token(token):
                has_code = True
            elif self._looks_like_text(token):
                has_text = True
        return has_code and has_text

    def generate_translated_workbook(
        self,
        original_workbook: OpenpyxlWorkbook,
        temp_path: Path,
        tamil_translations: Dict[str, str],
        hindi_translations: Dict[str, str],
        cell_mapping: List[Dict[str, Any]],
        output_path: Path
    ) -> None:
        """
        Generate a translated Excel workbook with Tamil and Hindi sheets.

        The output holds every English sheet followed by one copy per language
        (Tamil_<sheet>, Hindi_<sheet>). Sheets of the upload that are already in
        Tamil or Hindi are left out. The copies are made
        inside the .xlsx file of temp_path, so the template (logo, pictures,
        shapes, borders, merged cells, print setup) stays exactly as it is and
        only the translated cell texts change.
        """
        logger.info("Creating translated workbook")

        source_titles = [sheet.title for sheet in self._source_sheets(original_workbook)]
        used_titles = set(source_titles)
        # The English sheets come first, exactly as they were uploaded
        copies: List[SheetCopy] = [
            SheetCopy(title, title, {}, {title: title}) for title in source_titles
        ]

        languages = [("Tamil", tamil_translations), ("Hindi", hindi_translations)]
        for language, translations in languages:
            # All titles of a language are needed first, so formulas that point to
            # another sheet can be redirected to the copy in the same language.
            title_map: Dict[str, str] = {}
            for title in source_titles:
                title_map[title] = self._make_sheet_title(used_titles, language, title)
                used_titles.add(title_map[title])

            for title in source_titles:
                cell_values = {
                    f"{get_column_letter(mapping['column'])}{mapping['row']}": self._translate_cell_value(
                        mapping["original_value"], mapping["translatable_parts"], translations
                    )
                    for mapping in cell_mapping if mapping["sheet_name"] == title
                }
                copies.append(SheetCopy(title, title_map[title], cell_values, title_map))

        write_sheet_copies(temp_path, output_path, copies)
        logger.info("Translated workbook saved", filename=str(output_path))

    def _make_sheet_title(self, used_titles: Set[str], prefix: str, title: str) -> str:
        """Build a unique sheet title with language prefix, within Excel's length limit."""
        base_title = prefix if title == "Sheet" else f"{prefix}_{title}"
        new_title = base_title[:MAX_SHEET_TITLE_LENGTH]

        counter = 2
        while new_title in used_titles:
            suffix = f"_{counter}"
            new_title = base_title[:MAX_SHEET_TITLE_LENGTH - len(suffix)] + suffix
            counter += 1
        return new_title

    def _translate_cell_value(
        self,
        original_value: str,
        translatable_parts: List[str],
        translation_dict: Dict[str, str]
    ) -> str:
        """Translate a cell value while preserving protected values."""
        if not translatable_parts:
            return original_value

        result = original_value
        position = 0
        for part in translatable_parts:
            index = result.find(part, position)
            if index == -1:
                continue

            translated = translation_dict.get(part)
            if not translated or not self._numbers_preserved(part, translated):
                # Never risk a changed quantity, price or date in a PO: keep English
                if translated:
                    logger.warning("Translation changed a number, original text kept", text=part)
                position = index + len(part)
                continue

            result = result[:index] + translated + result[index + len(part):]
            position = index + len(translated)
        return result

    def _numbers_preserved(self, original: str, translated: str) -> bool:
        """Check that a translation contains exactly the numbers of the original text."""
        return sorted(re.findall(r'\d+', original)) == sorted(re.findall(r'\d+', translated))

    def get_workbook_info(self, file_path: Path) -> Dict[str, Any]:
        """Get information about an Excel workbook."""
        workbook = load_workbook(file_path, data_only=True)
        sheet_info = []
        for sheet in workbook.worksheets:
            sheet_info.append({
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column
            })
        return {
            "filename": Path(file_path).name,
            "sheet_count": len(workbook.worksheets),
            "sheets": sheet_info
        }

    def create_sample_po(self, file_path: Path) -> None:
        """Create a sample PO file for testing."""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Purchase Order"

        headers = [
            ("A1", "Purchase Order"),
            ("A2", "PO Number"),
            ("B2", "PO-10025"),
            ("A3", "Supplier Name"),
            ("B3", "ABC Manufacturing Co."),
            ("A4", "Material Description"),
            ("B4", "MAT-1001 Hydraulic Pump"),
            ("A5", "Quantity"),
            ("B5", 100),
            ("A6", "Unit Price"),
            ("B6", 50.00),
            ("A7", "Delivery Address"),
            ("B7", "123 Industrial Blvd, City, State 12345"),
            ("A8", "Payment Terms"),
            ("B8", "Net 30 Days"),
            ("A9", "Delivery Date"),
            ("B9", "07/10/2026"),
            ("A10", "Part Number"),
            ("B10", "JCB-12345"),
            ("A11", "Total Amount"),
            ("B11", "=B5*B6"),
            ("A12", "Contact"),
            ("B12", "purchase@example.com"),
            ("A13", "Remarks"),
            ("B13", "Deliver before noon.\nHandle with care."),
        ]

        for cell_ref, value in headers:
            sheet[cell_ref] = value

        sheet.merge_cells("A1:B1")
        sheet["A1"].font = Font(bold=True, size=14, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", start_color="1F4E78")
        sheet["A1"].alignment = Alignment(horizontal="center")
        for row in range(2, 14):
            sheet.cell(row=row, column=1).font = Font(bold=True)
        sheet["B6"].number_format = '"$"#,##0.00'
        sheet["B11"].number_format = '"$"#,##0.00'
        sheet["B13"].alignment = Alignment(wrap_text=True)
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 42
        sheet.freeze_panes = "A2"

        terms = workbook.create_sheet("Terms")
        terms["A1"] = "Terms and Conditions"
        terms["A1"].font = Font(bold=True)
        terms["A2"] = "Goods must match the specification of PO-10025."
        terms["A3"] = "Payment will be released after quality inspection."
        terms["A4"] = "Order value"
        terms["B4"] = "='Purchase Order'!B11"
        terms.column_dimensions["A"].width = 55

        workbook.save(file_path)
