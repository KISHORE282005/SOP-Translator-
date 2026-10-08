"""End-to-end tests for the upload, translate and download endpoints."""

import io
import uuid

from openpyxl import Workbook, load_workbook

from app.config import settings

XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def upload(client, name, content):
    return client.post("/api/upload", files={"file": (name, content, XLSX_MIME_TYPE)})


def test_root_and_health(client):
    assert client.get("/").status_code == 200
    assert client.get("/health").json()["status"] == "ok"


def test_upload_translate_download(client, sample_po):
    response = upload(client, "Sample PO.xlsx", sample_po.read_bytes())
    assert response.status_code == 200
    body = response.json()
    assert body["filename"] == "Sample PO.xlsx"
    assert body["status"] == "uploaded"
    file_id = body["file_id"]

    response = client.post(f"/api/translate/{file_id}")
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "completed"
    assert body["languages"] == ["Tamil", "Hindi"]
    assert body["output_filename"] == "Sample_PO_Translated.xlsx"
    assert body["translated_cells"] > 0

    response = client.get(f"/api/download/{file_id}")
    assert response.status_code == 200
    assert "Sample_PO_Translated.xlsx" in response.headers["content-disposition"]

    output = load_workbook(io.BytesIO(response.content))
    assert output["Tamil_Purchase Order"]["A3"].value == "[ta] Supplier Name"
    assert output["Hindi_Purchase Order"]["A3"].value == "[hi] Supplier Name"
    assert output["Tamil_Purchase Order"]["B2"].value == "PO-10025"


def test_translate_can_be_repeated(client, sample_po):
    file_id = upload(client, "po.xlsx", sample_po.read_bytes()).json()["file_id"]
    assert client.post(f"/api/translate/{file_id}").status_code == 200
    assert client.post(f"/api/translate/{file_id}").status_code == 200
    assert client.get(f"/api/download/{file_id}").status_code == 200


def test_upload_rejects_wrong_extension(client):
    response = upload(client, "po.csv", b"a,b,c")
    assert response.status_code == 400


def test_upload_rejects_invalid_workbook(client):
    response = upload(client, "po.xlsx", b"this is not an excel file")
    assert response.status_code == 400


def test_upload_rejects_empty_workbook(client):
    buffer = io.BytesIO()
    Workbook().save(buffer)
    response = upload(client, "empty.xlsx", buffer.getvalue())
    assert response.status_code == 400


def test_upload_rejects_large_file(client, monkeypatch):
    monkeypatch.setattr(settings, "max_file_size", 10)
    response = upload(client, "po.xlsx", b"x" * 100)
    assert response.status_code == 413


def test_upload_does_not_escape_upload_folder(client, sample_po):
    response = upload(client, "../../evil.xlsx", sample_po.read_bytes())
    assert response.status_code == 200
    file_id = response.json()["file_id"]
    stored = list(settings.upload_dir.glob(f"{file_id}_*"))
    assert [path.name for path in stored] == [f"{file_id}_evil.xlsx"]


def test_translate_unknown_file(client):
    assert client.post(f"/api/translate/{uuid.uuid4()}").status_code == 404


def test_translate_rejects_wildcard_file_id(client, sample_po):
    upload(client, "po.xlsx", sample_po.read_bytes())
    assert client.post("/api/translate/*").status_code == 404


def test_translate_workbook_without_text(client):
    workbook = Workbook()
    workbook.active["A1"] = 100
    workbook.active["A2"] = "PO-10025"
    buffer = io.BytesIO()
    workbook.save(buffer)

    file_id = upload(client, "numbers.xlsx", buffer.getvalue()).json()["file_id"]
    assert client.post(f"/api/translate/{file_id}").status_code == 400


def test_download_before_translation(client, sample_po):
    file_id = upload(client, "po.xlsx", sample_po.read_bytes()).json()["file_id"]
    assert client.get(f"/api/download/{file_id}").status_code == 404


def test_download_rejects_wildcard_file_id(client, sample_po):
    file_id = upload(client, "po.xlsx", sample_po.read_bytes()).json()["file_id"]
    client.post(f"/api/translate/{file_id}")
    assert client.get("/api/download/*").status_code == 404
