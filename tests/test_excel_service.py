"""Tests for the Excel processing service."""

import pytest
from openpyxl import Workbook, load_workbook


def fake_translations(texts, code):
    return {text: f"[{code}] {text}" for text in texts}


def translate_sample(excel_service, source_path, output_path):
    workbook = excel_service.read_workbook(source_path)
    content = excel_service.identify_translatable_content(workbook)
    excel_service.generate_translated_workbook(
        workbook,
        source_path,
        fake_translations(content["english_texts"], "ta"),
        fake_translations(content["english_texts"], "hi"),
        content["cell_mapping"],
        output_path,
    )
    return content, load_workbook(output_path)


@pytest.mark.parametrize("value", [
    "PO-10025", "PO/2026/001", "MAT-1001", "JCB-12345", "EMP001", "V002",
    "100", "5,000.00", "15%", "07/10/2026", "2026-07-10", "$50.00",
    "https://example.com/po", "purchase@example.com", "M8x40", "=B5*B6",
])
def test_protected_values_are_not_translated(excel_service, value):
    assert excel_service._is_protected_value(value) is True


@pytest.mark.parametrize("value", [
    "Purchase Order", "Supplier Name", "Net 30 Days", "ABC Manufacturing Co.",
])
def test_text_values_are_not_protected(excel_service, value):
    assert excel_service._is_protected_value(value) is False


@pytest.mark.parametrize("value, expected", [
    ("Hydraulic Pump", ["Hydraulic Pump"]),
    ("MAT-1001 Hydraulic Pump", ["Hydraulic Pump"]),
    ("Pump (MAT-1001) for JCB-12345 loader", ["Pump", "for", "loader"]),
    ("Net 30 Days", ["Net 30 Days"]),
    ("Deliver before noon.\nHandle with care.", ["Deliver before noon.", "Handle with care."]),
    ("PO-10025 / MAT-1001", []),
    ("12 x 40", []),
])
def test_extract_translatable_parts(excel_service, value, expected):
    assert excel_service._extract_translatable_parts(value) == expected


def test_translate_cell_value_keeps_codes(excel_service):
    result = excel_service._translate_cell_value(
        "MAT-1001 Hydraulic Pump", ["Hydraulic Pump"], {"Hydraulic Pump": "ஹைட்ராலிக் பம்ப்"}
    )
    assert result == "MAT-1001 ஹைட்ராலிக் பம்ப்"


def test_translate_cell_value_keeps_english_when_a_number_changes(excel_service):
    result = excel_service._translate_cell_value(
        "Net 30 Days", ["Net 30 Days"], {"Net 30 Days": "நிகர 3 நாட்கள்"}
    )
    assert result == "Net 30 Days"


def test_identify_translatable_content(excel_service, sample_po):
    workbook = excel_service.read_workbook(sample_po)
    content = excel_service.identify_translatable_content(workbook)

    texts = content["english_texts"]
    assert "Purchase Order" in texts
    assert "Hydraulic Pump" in texts
    assert "Net 30 Days" in texts
    assert len(texts) == len(set(texts))
    for protected in ("PO-10025", "JCB-12345", "07/10/2026", "purchase@example.com", "=B5*B6"):
        assert protected not in texts

    mixed = [m for m in content["cell_mapping"] if m["original_value"] == "MAT-1001 Hydraulic Pump"]
    assert len(mixed) == 1
    assert mixed[0]["is_mixed"] is True


def test_generated_workbook_has_tamil_and_hindi_sheets(excel_service, sample_po, tmp_path):
    _, output = translate_sample(excel_service, sample_po, tmp_path / "out.xlsx")

    assert output.sheetnames == [
        "Purchase Order", "Terms",
        "Tamil_Purchase Order", "Tamil_Terms", "Hindi_Purchase Order", "Hindi_Terms",
    ]
    assert output["Purchase Order"]["A1"].value == "Purchase Order"
    assert output["Terms"]["B4"].value == "='Purchase Order'!B11"

    tamil = output["Tamil_Purchase Order"]
    hindi = output["Hindi_Purchase Order"]
    assert tamil["A1"].value == "[ta] Purchase Order"
    assert hindi["A1"].value == "[hi] Purchase Order"
    assert tamil["B4"].value == "MAT-1001 [ta] Hydraulic Pump"
    assert hindi["B13"].value == "[hi] Deliver before noon.\n[hi] Handle with care."


def test_generated_workbook_preserves_protected_values(excel_service, sample_po, tmp_path):
    _, output = translate_sample(excel_service, sample_po, tmp_path / "out.xlsx")

    for sheet_name in ("Tamil_Purchase Order", "Hindi_Purchase Order"):
        sheet = output[sheet_name]
        assert sheet["B2"].value == "PO-10025"
        assert sheet["B5"].value == 100
        assert sheet["B6"].value == 50.00
        assert sheet["B9"].value == "07/10/2026"
        assert sheet["B10"].value == "JCB-12345"
        assert sheet["B11"].value == "=B5*B6"
        assert sheet["B12"].value == "purchase@example.com"

    # Codes inside a sentence stay as they are
    assert "PO-10025" in output["Tamil_Terms"]["A2"].value


def test_generated_workbook_preserves_formatting(excel_service, sample_po, tmp_path):
    _, output = translate_sample(excel_service, sample_po, tmp_path / "out.xlsx")

    for sheet_name in ("Tamil_Purchase Order", "Hindi_Purchase Order"):
        sheet = output[sheet_name]
        assert "A1:B1" in [str(merged) for merged in sheet.merged_cells.ranges]
        assert sheet["A1"].font.bold is True
        assert sheet["A1"].fill.start_color.rgb.endswith("1F4E78")
        assert sheet["B6"].number_format == '"$"#,##0.00'
        assert sheet.column_dimensions["B"].width == 42
        assert sheet.freeze_panes == "A2"


def test_cross_sheet_formulas_point_to_the_same_language(excel_service, sample_po, tmp_path):
    _, output = translate_sample(excel_service, sample_po, tmp_path / "out.xlsx")

    assert output["Tamil_Terms"]["B4"].value == "='Tamil_Purchase Order'!B11"
    assert output["Hindi_Terms"]["B4"].value == "='Hindi_Purchase Order'!B11"


def test_source_workbook_is_not_modified(excel_service, sample_po, tmp_path):
    before = sample_po.read_bytes()
    translate_sample(excel_service, sample_po, tmp_path / "out.xlsx")
    assert sample_po.read_bytes() == before


def test_long_sheet_titles_stay_unique_and_within_excel_limit(excel_service, tmp_path):
    source_path = tmp_path / "long.xlsx"
    workbook = Workbook()
    workbook.active.title = "X" * 30 + "1"
    workbook.active["A1"] = "Purchase Order"
    second = workbook.create_sheet("X" * 30 + "2")
    second["A1"] = "Delivery Address"
    workbook.save(source_path)

    _, output = translate_sample(excel_service, source_path, tmp_path / "out.xlsx")

    assert len(output.sheetnames) == 6
    assert len(set(output.sheetnames)) == 6
    assert all(len(title) <= 31 for title in output.sheetnames)


def test_images_are_copied_to_every_language(excel_service, sample_po, tmp_path):
    PILImage = pytest.importorskip("PIL.Image")
    from openpyxl.drawing.image import Image

    logo_path = tmp_path / "logo.png"
    PILImage.new("RGB", (80, 40), "navy").save(logo_path)
    workbook = load_workbook(sample_po)
    workbook["Purchase Order"].add_image(Image(str(logo_path)), "D2")
    workbook.save(sample_po)

    _, output = translate_sample(excel_service, sample_po, tmp_path / "out.xlsx")

    assert len(output["Purchase Order"]._images) == 1
    assert len(output["Tamil_Purchase Order"]._images) == 1
    assert len(output["Hindi_Purchase Order"]._images) == 1


def sop_workbook(path):
    """A small workbook laid out like the SOP template: header, steps, footer."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SOP PAGE1"
    sheet["A1"] = "STANDARD OPERATING PROCEDURE"
    sheet["A2"] = "Seq.\nNo."
    sheet["B2"] = "Operation Sequence"
    sheet["A3"] = "1"
    sheet["B3"] = "Place the ARM_SUPPORT_PLATE on the fixture"
    sheet["A4"] = 2
    sheet["B4"] = "Wear safety gloves"
    sheet["A6"] = "Process Sheet Revision History"
    sheet["B7"] = "Safety Shoes"
    hindi = workbook.create_sheet("SOP PAGE1 (Hindi)")
    hindi["A1"] = "STANDARD OPERATING PROCEDURE"
    hindi["B3"] = "सही PPE पहनें"
    workbook.save(path)
    return path


def test_sop_sheet_is_translated_completely(excel_service, tmp_path):
    source = sop_workbook(tmp_path / "sop.xlsx")
    _, output = translate_sample(excel_service, source, tmp_path / "out.xlsx")

    # The sheet that is already in Hindi is not used as a source
    assert output.sheetnames == ["SOP PAGE1", "Tamil_SOP PAGE1", "Hindi_SOP PAGE1"]

    english = output["SOP PAGE1"]
    assert english["A1"].value == "STANDARD OPERATING PROCEDURE"
    assert english["B4"].value == "Wear safety gloves"

    tamil = output["Tamil_SOP PAGE1"]
    assert tamil["A1"].value == "[ta] STANDARD OPERATING PROCEDURE"
    assert tamil["B2"].value == "[ta] Operation Sequence"
    assert tamil["B3"].value == "[ta] Place the ARM_SUPPORT_PLATE [ta] on the fixture"
    assert tamil["A6"].value == "[ta] Process Sheet Revision History"
    assert tamil["B7"].value == "[ta] Safety Shoes"
    assert output["Hindi_SOP PAGE1"]["B7"].value == "[hi] Safety Shoes"


def test_shapes_and_tables_survive_the_copy(excel_service, sample_po, tmp_path):
    import zipfile
    from openpyxl.worksheet.table import Table

    workbook = load_workbook(sample_po)
    workbook["Purchase Order"].add_table(Table(displayName="Items", ref="A2:B5"))
    workbook.save(sample_po)

    _, output = translate_sample(excel_service, sample_po, tmp_path / "out.xlsx")

    names = [table for sheet in output.worksheets for table in sheet.tables]
    assert len(names) == 3 and len(set(names)) == 3
    with zipfile.ZipFile(tmp_path / "out.xlsx") as archive:
        assert archive.testzip() is None
